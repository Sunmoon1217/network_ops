"""LTM 主备合并接口（``/api/analysis/ltm-groups/*``）测试。

覆盖：① 主备组逐层合并（VS→池→成员并集、device 逗号拼接）；② 标量冲突取主机名第一台
并标注；③ 单机/堆叠不合并；④ 组/设备过滤、搜索、排序、分页；⑤ facets 与导出。
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from assets.models import Device, DeviceGroup, DeviceGroupMember, LtmPool, LtmPoolMember, LtmVirtualServer

BASE = "/api/analysis/ltm-groups/"


@pytest.fixture
def api():
    return APIClient()


def _device(hostname: str) -> Device:
    return Device.objects.create(hostname=hostname, device_type="slb")


def _group(name: str, group_type: str, devices: list[Device]) -> DeviceGroup:
    group = DeviceGroup.objects.create(name=name, group_type=group_type)
    for device in devices:
        DeviceGroupMember.objects.create(group=group, device=device)
    return group


def _seed_vs(device, name="/Common/vs_web", pool="pool_web", **kwargs) -> LtmVirtualServer:
    defaults = {"vs_address": "10.0.0.100", "vs_port": "443", "protocol": "tcp"}
    defaults.update(kwargs)
    return LtmVirtualServer.objects.create(device=device, name=name, pool=pool, **defaults)


def _seed_pool(device, name="pool_web", mode="round-robin", monitors=None) -> LtmPool:
    return LtmPool.objects.create(device=device, name=name, mode=mode, monitors=monitors or ["http"])


def _seed_member(device, pool_name="pool_web", name="/Common/node_a", address="10.1.1.1", port="80"):
    return LtmPoolMember.objects.create(
        device=device, pool_name=pool_name, name=name, address=address, port=port
    )


def _read_xlsx(resp):
    return load_workbook(BytesIO(resp.content)).active


@pytest.mark.django_db
def test_ha_pair_merges_vs_pool_and_members(api):
    """主备两台 SLB 的同一 VS 合并成一行，成员取并集、device 逗号拼接。"""
    a, b = _device("slb-a"), _device("slb-b")
    _group("slb-pair", "ha", [a, b])
    for device in (a, b):
        _seed_vs(device)
        _seed_pool(device)
    _seed_member(a, name="/Common/node_a")
    _seed_member(b, name="/Common/node_a")
    _seed_member(b, name="/Common/node_b", address="10.1.1.2")

    rows = api.get(BASE).data["results"]
    assert len(rows) == 1
    row = rows[0]
    assert row["group_name"] == "slb-pair"
    assert row["device"] == "slb-a,slb-b"
    assert row["name"] == "/Common/vs_web"
    assert row["inconsistent"] is False

    assert len(row["pools"]) == 1
    pool = row["pools"][0]
    assert pool["name"] == "pool_web"
    assert pool["device"] == "slb-a,slb-b"
    assert pool["mode"] == "round-robin"
    # 成员并集：node_a 两边都有、node_b 只在 b 上
    members = {member["name"]: member for member in pool["members"]}
    assert set(members) == {"/Common/node_a", "/Common/node_b"}
    assert members["/Common/node_a"]["device"] == "slb-a,slb-b"
    assert members["/Common/node_b"]["device"] == "slb-b"


@pytest.mark.django_db
def test_scalar_conflict_takes_first_hostname_and_flags(api):
    """VS 地址/池模式/成员地址在两台不同时：取主机名第一台并逐字段标注。"""
    a, b = _device("slb-a"), _device("slb-b")
    _group("slb-pair", "ha", [a, b])
    # b 先建，验证最终取的是主机名第一台而不是插入顺序
    _seed_vs(b, vs_address="10.9.9.9")
    _seed_vs(a, vs_address="10.0.0.100")
    _seed_pool(b, mode="least-conn")
    _seed_pool(a, mode="round-robin")
    _seed_member(b, name="/Common/node_a", address="10.1.1.99")
    _seed_member(a, name="/Common/node_a", address="10.1.1.1")

    row = api.get(BASE).data["results"][0]
    assert row["vs_address"] == "10.0.0.100"
    assert row["inconsistent"] is True
    assert row["conflict_fields"] == ["vs_address"]

    pool = row["pools"][0]
    assert pool["mode"] == "round-robin"
    assert "mode" in pool["conflict_fields"]

    member = pool["members"][0]
    assert member["address"] == "10.1.1.1"
    assert member["inconsistent"] is True
    assert member["conflict_fields"] == ["address"]


@pytest.mark.django_db
def test_single_and_stack_groups_are_not_merged(api):
    """单机 / 堆叠组不参与合并：成员各自成为单机组，各出一行。"""
    a, b = _device("slb-a"), _device("slb-b")
    _group("slb-stack", "stack", [a, b])
    for device in (a, b):
        _seed_vs(device)
        _seed_pool(device)

    rows = api.get(BASE).data["results"]
    assert [(row["group_name"], row["device"]) for row in rows] == [("slb-a", "slb-a"), ("slb-b", "slb-b")]


@pytest.mark.django_db
def test_group_and_device_filter_select_whole_group(api):
    """?group= 组键/组名、?device= 组内任一设备，都返回整组合并行。"""
    a, b, other = _device("slb-a"), _device("slb-b"), _device("slb-other")
    group = _group("slb-pair", "failover", [a, b])
    for device in (a, b, other):
        _seed_vs(device)
        _seed_pool(device)

    for params in ({"group": f"group-{group.pk}"}, {"group": "slb-pair"}, {"device": b.pk}):
        results = api.get(BASE, params).data["results"]
        assert [row["device"] for row in results] == ["slb-a,slb-b"], params

    assert [row["device"] for row in api.get(BASE, {"device": other.pk}).data["results"]] == ["slb-other"]
    assert api.get(BASE, {"device": "abc"}).data["count"] == 0
    assert api.get(BASE, {"group": "nope"}).data["count"] == 0


@pytest.mark.django_db
def test_protocol_status_search_ordering_and_pagination(api):
    device = _device("slb-a")
    _seed_vs(device, name="vs_b", pool="pool_b", vs_address="10.0.0.2", protocol="tcp", status="enabled")
    _seed_vs(device, name="vs_a", pool="pool_a", vs_address="10.0.0.1", protocol="udp", status="disabled")
    _seed_vs(device, name="vs_c", pool="pool_c", vs_address="10.0.0.3")
    for pool_name in ("pool_a", "pool_b", "pool_c"):
        _seed_pool(device, name=pool_name)

    def names(params):
        return [row["name"] for row in api.get(BASE, params).data["results"]]

    assert names({"protocol": "udp"}) == ["vs_a"]
    assert names({"status": "disabled"}) == ["vs_a"]
    assert names({"search": "10.0.0.2"}) == ["vs_b"]
    assert names({"search": "pool_b"}) == ["vs_b"]
    assert names({"ordering": "-name"}) == ["vs_c", "vs_b", "vs_a"]
    payload = api.get(BASE, {"page_size": 2}).data
    assert payload["count"] == 3 and len(payload["results"]) == 2 and payload["next"]


@pytest.mark.django_db
def test_facets_include_groups_protocols_and_statuses(api):
    a, b = _device("slb-a"), _device("slb-b")
    _group("slb-pair", "ha", [a, b])
    for device in (a, b):
        _seed_vs(device, protocol="tcp", status="enabled")
        _seed_pool(device)

    facets = api.get(f"{BASE}facets/").data
    assert [group["name"] for group in facets["groups"]] == ["slb-pair"]
    assert facets["groups"][0]["devices"] == ["slb-a", "slb-b"]
    assert facets["protocols"] == ["tcp"]
    assert facets["statuses"] == ["enabled"]


@pytest.mark.django_db
def test_export_flat_rows_show_joined_device_and_conflict(api):
    from analysis.api.ltm_groups import EXPORT_COLUMN_WIDTHS, EXPORT_HEADERS

    a, b = _device("slb-a"), _device("slb-b")
    _group("slb-pair", "ha", [a, b])
    for device in (a, b):
        _seed_vs(device)
        _seed_pool(device)
    _seed_member(a, name="/Common/node_a")
    _seed_member(b, name="/Common/node_b", address="10.1.1.2")
    # 只在 a 上存在的 VS → 不冲突（条目集合的并集语义）
    _seed_vs(a, name="vs_only_a", vs_address="10.0.0.201", pool="pool_a")
    _seed_pool(a, name="pool_a")

    resp = api.get(f"{BASE}export/")
    assert resp.status_code == 200
    assert len(EXPORT_COLUMN_WIDTHS) == len(EXPORT_HEADERS)
    sheet = _read_xlsx(resp)
    assert [cell.value for cell in sheet[1]] == EXPORT_HEADERS
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]

    merged = [row for row in rows if row[2] == "/Common/vs_web"]
    assert {row[11] for row in merged} == {"10.1.1.1#80", "10.1.1.2#80"}
    assert {row[0] for row in merged} == {"slb-a,slb-b"}
    assert {row[12] for row in merged} == {None}  # 无冲突（xlsx 空格读回 None）

    only_a = [row for row in rows if row[2] == "vs_only_a"][0]
    assert only_a[0] == "slb-a"
