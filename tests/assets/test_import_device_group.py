"""Excel 导入：「设备组」Sheet 与模板下载接口。

设备组 Sheet 是 **一行 = 一个设备组**（组类型 | 设备名(多个用逗号分隔) | 描述），
组名与角色都不填表、由导入推：

- 非集群：组名 = 主设备（第一台）hostname，第一台 = 主、第二台 = 备
- 集群：组名 = 所有设备名的最长公共前缀（``cs-1``/``cs-2`` → ``cs-``），组内全部 = 成员
- 行里的设备必须都已存在，缺任一台**整行报错、不建半截组**
- 行是该组的**完整快照**：不在这一行里的既有成员会被摘掉

另覆盖模板接口：Sheet 与列头取自 `IMPORT_SHEETS`，与 `import_excel` 分发的是同一份表
（改列头只改一处，模板与导入不可能漂移）。
"""

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from assets.device_import import IMPORT_SHEETS, _device_roles, _group_name, _import_device_groups
from assets.models import Device, DeviceGroup

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _device(hostname: str) -> Device:
    return Device.objects.create(hostname=hostname, device_type="switch")


def _sheet(rows, headers=None):
    ws = Workbook().active
    ws.title = "设备组"  # 直接调 importer 无所谓，走 HTTP 分发时 Sheet 名就是 key
    ws.append(headers or IMPORT_SHEETS["设备组"]["headers"])
    for row in rows:
        ws.append(row)
    return ws


def _members(group):
    return dict(group.members.values_list("device__hostname", "device_role"))


# ---------------------------------------------------------------------------
# 组名 / 角色的推导规则（纯函数，不碰库）
# ---------------------------------------------------------------------------


def test_group_name_rules():
    assert _group_name("cluster", ["cs-1", "cs-2"]) == "cs-", "集群组名 = 设备名的逐字符交集"
    assert _group_name("cluster", ["cs-1"]) == "cs-1", "单台集群组名就是它自己"
    assert _group_name("cluster", ["rt1", "sw2"]) == "rt1", "毫无公共前缀时回退第一台"
    assert _group_name("ha", ["core-sw", "core-sw-b"]) == "core-sw", "非集群组名 = 主设备 hostname"


def test_device_role_rules():
    assert _device_roles("ha", 2) == ["master", "backup"]
    assert _device_roles("stack", 1) == ["master"]
    assert _device_roles("cluster", 3) == ["member", "member", "member"]
    assert _device_roles("failover", 3) == ["master", "backup", "member"], "非集群第三台起按成员"


# ---------------------------------------------------------------------------
# 导入行为
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_one_row_creates_one_group_with_derived_name_and_roles():
    """一行 = 一个组；组名取第一台（主设备）hostname，角色按位置推"""
    _device("imp-ha-m")
    _device("imp-ha-b")

    result = _import_device_groups(_sheet([["主备", "imp-ha-m, imp-ha-b", "出口主备"]]))

    assert result == {"created": 1, "updated": 0, "errors": []}, result
    group = DeviceGroup.objects.get()
    assert group.name == "imp-ha-m"
    assert group.group_type == "ha"
    assert group.description == "出口主备"
    assert _members(group) == {"imp-ha-m": "master", "imp-ha-b": "backup"}


@pytest.mark.django_db
def test_cluster_group_name_is_common_prefix():
    """集群：一个格子里逗号分隔多台，组名 = 所有设备名的交集（cs-1/cs-2 → cs-）"""
    _device("cs-1")
    _device("cs-2")

    result = _import_device_groups(_sheet([["集群", "cs-1,cs-2", ""]]))

    assert result == {"created": 1, "updated": 0, "errors": []}, result
    group = DeviceGroup.objects.get()
    assert group.name == "cs-"
    assert group.group_type == "cluster"
    assert _members(group) == {"cs-1": "member", "cs-2": "member"}


@pytest.mark.django_db
def test_cluster_without_common_prefix_falls_back_to_first_device():
    """设备名毫无公共前缀时组名回退第一台，不能出现空组名"""
    _device("rt-x")
    _device("sw-y")

    _import_device_groups(_sheet([["集群", "rt-x、sw-y", ""]]))  # 中文顿号也认

    group = DeviceGroup.objects.get()
    assert group.name == "rt-x"
    assert set(group.members.values_list("device__hostname", flat=True)) == {"rt-x", "sw-y"}


@pytest.mark.django_db
def test_blank_type_defaults_to_single_and_first_device_is_master():
    """组类型留空 → single；单台成组时第一台即主"""
    device = _device("imp-solo")

    result = _import_device_groups(_sheet([["", "imp-solo", ""]]))

    assert result["created"] == 1 and not result["errors"], result
    group = DeviceGroup.objects.get()
    assert group.name == "imp-solo"
    assert group.group_type == "single"
    assert _members(group) == {device.hostname: "master"}


@pytest.mark.django_db
def test_bad_rows_reported_without_partial_group():
    """未知设备 / 组类型 / 空设备名逐行报错；报错行不许留下半截组"""
    _device("imp-ok")

    result = _import_device_groups(
        _sheet(
            [
                ["主备", "imp-not-exist", ""],
                ["机架", "imp-ok", ""],
                ["ha", "", ""],
            ]
        )
    )

    assert result["created"] == 0
    assert len(result["errors"]) == 3, result
    assert any("设备 'imp-not-exist' 不存在" in e for e in result["errors"])
    assert any("未知组类型" in e for e in result["errors"])
    assert any("缺少设备名" in e for e in result["errors"])
    assert not DeviceGroup.objects.exists(), "报错行不该建出组"


@pytest.mark.django_db
def test_reimport_counts_as_update_and_row_is_full_snapshot():
    """重复导入按「更新」计数；行里没有的既有成员会被摘掉（组成以本次为准）"""
    _device("imp-snap-a")
    _device("imp-snap-b")

    first = _import_device_groups(_sheet([["主备", "imp-snap-a,imp-snap-b", ""]]))
    second = _import_device_groups(_sheet([["主备", "imp-snap-a", ""]]))

    assert first["created"] == 1 and first["updated"] == 0, first
    assert second["created"] == 0 and second["updated"] == 1, second

    group = DeviceGroup.objects.get(name="imp-snap-a")
    assert _members(group) == {"imp-snap-a": "master"}, "不在行里的成员要被摘掉"


# ---------------------------------------------------------------------------
# 模板与 HTTP 分发
# ---------------------------------------------------------------------------


def test_template_lists_every_import_sheet():
    """模板的 Sheet 与列头 == IMPORT_SHEETS（同一份表），设备组 Sheet 必须在"""
    resp = APIClient().get("/api/assets/import-template/")

    assert resp.status_code == 200
    assert "spreadsheetml" in resp["Content-Type"]

    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == list(IMPORT_SHEETS), "模板 Sheet 清单/顺序要与导入分发一致"
    assert "设备组" in wb.sheetnames

    for sheet_name, spec in IMPORT_SHEETS.items():
        header_row = [cell.value for cell in wb[sheet_name][1]]
        assert header_row == spec["headers"], f"{sheet_name} 列头与导入实现不一致"
    assert wb["设备组"][1][0].value == "组类型"
    assert wb["设备组"][1][1].value == "设备名(多个用逗号分隔)"


@pytest.mark.django_db
def test_import_excel_dispatches_device_group_sheet():
    """走完整 HTTP 导入链：设备组 Sheet 被 IMPORT_SHEETS 认出并分发"""
    _device("imp-http-dev")

    buf = BytesIO()
    _sheet([["集群", "imp-http-dev", "集群组"]]).parent.save(buf)
    upload = SimpleUploadedFile("import.xlsx", buf.getvalue(), XLSX_CONTENT_TYPE)

    payload = APIClient().post("/api/assets/import-devices/", {"file": upload}).json()

    assert payload["success"] is True
    assert payload["results"]["设备组"]["created"] == 1, payload
    assert payload["results"]["配置文件"]["skipped"] is True, "未提供的 Sheet 应记为跳过"
    assert DeviceGroup.objects.filter(name="imp-http-dev", group_type="cluster").exists()
