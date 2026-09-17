"""互联网资产分析扁平表格与 xlsx 导出测试。

表格形态是「一行一条从 WideIP 到最终后端的完整链路」，各层级横向铺成列，
不再是带缩进的树形。导出用后端的 openpyxl 生成，前端不引入 xlsx 依赖。
"""

from io import BytesIO

import pytest
from django.test import Client
from openpyxl import load_workbook

from assets.api.analysis import EXPORT_HEADERS, analyze_device, build_path_rows
from assets.models import (
    Device,
    GtmPool,
    GtmServer,
    GtmVServer,
    GtmWideip,
    LtmPool,
    LtmPoolMember,
    LtmVirtualServer,
)

EXPORT_URL = "/api/assets/internet-analysis/export/"


def _device(hostname: str, device_type: str = "gslb") -> Device:
    return Device.objects.create(hostname=hostname, device_type=device_type)


def _wideip(device: Device, name: str, pools: list[str]) -> GtmWideip:
    return GtmWideip.objects.create(device=device, name=name, rtype="A", lb_mode="round-robin", pools=pools)


def _gtm_pool(device: Device, name: str, members: list) -> GtmPool:
    return GtmPool.objects.create(device=device, name=name, members=members)


def _gtm_vserver(device: Device, server_name: str, vs_name: str, ip: str, port: str) -> GtmVServer:
    server, _ = GtmServer.objects.get_or_create(device=device, name=server_name)
    return GtmVServer.objects.create(device=device, server=server, name=vs_name, ip_address=ip, port=port)


def _ltm_virtual(device: Device, name: str, ip: str, port: str, pool: str = "") -> LtmVirtualServer:
    return LtmVirtualServer.objects.create(device=device, name=name, vs_address=ip, vs_port=port, pool=pool)


def _build_resolved_chain(hostname: str = "ia-exp") -> Device:
    """搭一条能完整解析到后端成员的链路"""
    gslb = _device(hostname)
    _wideip(gslb, "www.example.com", ["pool_web"])
    _gtm_pool(gslb, "pool_web", [{"server_name": "s1", "vs_name": "vs1"}])
    _gtm_vserver(gslb, "s1", "vs1", "10.1.1.1", "80")

    ltm = _device(f"{hostname}-ltm", "slb")
    _ltm_virtual(ltm, "vs_web", "10.1.1.1", "80", pool="pool_backend")
    LtmPool.objects.create(device=ltm, name="pool_backend", mode="http")
    LtmPoolMember.objects.create(pool_name="pool_backend", name="m1", address="10.9.9.9", port="8080")
    return gslb


# ---------- 扁平结构 ----------


@pytest.mark.django_db
def test_path_rows_carry_every_level_in_one_row():
    """一行就是一条完整链路，各层级横向铺在列里"""
    gslb = _build_resolved_chain("ia-flat")
    rows = build_path_rows(analyze_device(gslb))

    assert len(rows) == 1
    row = rows[0]
    assert row["wideip"] == "www.example.com"
    assert row["wideip_type"] == "A"
    assert row["pool"] == "pool_web"
    assert row["pool_found"] is True
    assert row["member"] == "s1:vs1"
    assert "vs1" in row["gtm_vserver"] and "10.1.1.1:80" in row["gtm_vserver"]
    assert row["ltm_chain"] == "vs_web"
    assert row["backend_member"] == "m1"
    assert row["final_address"] == "10.9.9.9:8080"
    assert row["status_code"] == "resolved"


@pytest.mark.django_db
def test_path_rows_are_not_tree_shaped():
    """扁平行不含 depth / kind 这类树形字段"""
    gslb = _build_resolved_chain("ia-notree")
    row = build_path_rows(analyze_device(gslb))[0]

    assert "depth" not in row
    assert "kind" not in row


@pytest.mark.django_db
def test_multiple_members_produce_multiple_rows():
    """同一池的多个成员各占一行"""
    gslb = _device("ia-multi")
    _wideip(gslb, "www.example.com", ["pool_web"])
    _gtm_pool(
        gslb,
        "pool_web",
        [
            {"server_name": "s1", "vs_name": "vs1", "member_order": 0},
            {"server_name": "s2", "vs_name": "vs2", "member_order": 1},
        ],
    )
    _gtm_vserver(gslb, "s1", "vs1", "10.1.1.1", "80")
    _gtm_vserver(gslb, "s2", "vs2", "10.1.1.2", "80")

    rows = build_path_rows(analyze_device(gslb))

    assert [row["member"] for row in rows] == ["s1:vs1", "s2:vs2"]
    assert {row["status_code"] for row in rows} == {"ltm_not_found"}


@pytest.mark.django_db
def test_empty_pool_still_produces_a_row():
    """池没有成员时也要留一行，否则这条记录会从表里消失"""
    gslb = _device("ia-empty")
    _wideip(gslb, "www.example.com", ["pool_missing"])

    rows = build_path_rows(analyze_device(gslb))

    assert len(rows) == 1
    assert rows[0]["pool"] == "pool_missing"
    assert rows[0]["pool_found"] is False
    assert rows[0]["member"] == ""


# ---------- xlsx 导出 ----------


@pytest.mark.django_db
def test_export_returns_xlsx_workbook():
    gslb = _build_resolved_chain("ia-xlsx")

    res = Client().get(EXPORT_URL, {"device": gslb.pk})

    assert res.status_code == 200
    assert res["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in res["Content-Disposition"]
    assert "ia-xlsx" in res["Content-Disposition"]

    sheet = load_workbook(BytesIO(res.content)).active
    assert sheet.title == "互联网资产分析"
    assert [cell.value for cell in sheet[1]] == EXPORT_HEADERS

    # 表头 + 一行链路
    assert sheet.max_row == 2
    values = [sheet.cell(row=2, column=index).value for index in range(1, len(EXPORT_HEADERS) + 1)]
    assert values[0] == "www.example.com"
    assert values[2] == "pool_web"
    assert values[3] == "已找到"
    assert values[8] == "10.9.9.9:8080"
    assert values[9] == "已解析到后端"


@pytest.mark.django_db
def test_export_requires_device_param():
    assert Client().get(EXPORT_URL).status_code == 400


@pytest.mark.django_db
def test_export_unknown_device_returns_404():
    assert Client().get(EXPORT_URL, {"device": 999999}).status_code == 404
