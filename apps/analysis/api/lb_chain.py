"""负载均衡 / 域名解析页面的关联链聚合接口（只读）。

两个页面各把一条关联链压进一行展示：

- 负载均衡：``LtmVirtualServer → LtmPool → LtmPoolMember``（VS / pool / member）
- 域名解析：``GtmWideip → GtmPool → GtmVServer``（wideip / pool / vs）

逐个 ViewSet 让前端再拼必然 N+1（每行两次查询），所以这里按页聚合、
固定条数查库返回整链。**不建辅助表**：单页默认 50 行只有个位数查询，
且配置入库是低频动作、数据随时可能被重新解析覆盖，实时查即可——
真出现慢查询再考虑缓存表，别提前背上失效同步的包袱。

URL 前缀 ``/api/lb-chain/`` 沿用 analysis「路由跟 view 走」的惯例挂进
``analysis/api/urls.py``；聚合口径（GTM 成员三种形态的兼容）复用
``analysis.api.analysis._parse_member_entry``，避免两处各写一份漂移。
"""

from datetime import date
from io import BytesIO
from urllib.parse import quote

from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from analysis.api.analysis import _join_ip_port, _parse_member_entry
from assets.models import GtmPool, GtmServer, GtmVServer, GtmWideip, LtmPool, LtmPoolMember, LtmVirtualServer
from netops.pagination import StandardPagination

# 模糊搜索反查出的 (device, pool) 组合上限：搜索词极短时可能命中全表，
# 组合 Q 是逐条 OR 展开的，必须封顶防止构造出巨型查询
MAX_SEARCH_PAIRS = 500

# 允许的排序字段（与 LTM/GTM 各 ViewSet 的 ordering_fields 对齐），
# 默认按 设备 + 名字 排，保证翻页稳定
ORDERABLE = {"name", "device_hostname", "created_at"}
DEFAULT_ORDER = ("device__hostname", "name")


def _device_filter(request, qs):
    device_id = request.query_params.get("device")
    return qs.filter(device_id=device_id) if device_id else qs


def _pairs_q(pairs, value_field: str) -> Q:
    """把 (device_id, 名字) 组合列表展开成 OR 查询。

    组合条件没法用单个 ``__in`` 表达（两个维度要成对匹配），只能逐对 OR；
    ``value_field`` 是第二个维度落到的字段名（如 ``pool`` / ``name``）。
    调用方负责用 ``MAX_SEARCH_PAIRS`` 封顶。
    """
    q = Q()
    for device_id, value in pairs:
        q |= Q(device_id=device_id, **{value_field: value})
    return q


def _order(qs, request, allowed_extra: dict):
    """按 ?ordering= 排序：前缀 '-' 即倒序，未知字段回退默认序。"""
    raw = request.query_params.get("ordering") or ""
    desc = raw.startswith("-")
    field = raw.lstrip("-")
    if field in ORDERABLE:
        field = allowed_extra.get(field, field)
        return qs.order_by(f"-{field}" if desc else field)
    return qs.order_by(*DEFAULT_ORDER)


def _page(request, qs):
    """统一走 StandardPagination（50/页、上限 500），返回 (当页列表, 分页器)。"""
    paginator = StandardPagination()
    page = paginator.paginate_queryset(qs, request)
    return list(page or []), paginator


# ---------------------------------------------------------------------------
# LTM：VS → 池 → 成员
# ---------------------------------------------------------------------------


def _split_addr_port(term: str):
    """把 ``地址:端口`` / ``地址#端口`` 拆成 (地址, 端口)；拆不出合法组合返回 None。

    支持的检索形态与坑：

    - ``#`` 是项目展示层的地址端口分隔符（与互联网资产分析同约定），天然无歧义；
    - ``:`` 必须防 IPv6 误拆——``2001:db8::1`` 的最后一段也是纯数字，直接按最后
      一个冒号切会把地址切残。规则：**最后一段是 1~5 位纯数字、且地址部分不以
      冒号结尾**才拆（``10.0.0.1:443``、``2001:db8::1:80`` 能拆，``2001:db8::1``
      不拆、退化成整串模糊匹配地址）；
    - 端口用**精确匹配**（搜 ``:80`` 不能命中 ``8080``）。
    """
    if "#" in term:
        addr, sep, port = term.rpartition("#")
    elif ":" in term:
        addr, sep, port = term.rpartition(":")
    else:
        return None
    if not (addr and sep and port.isdigit() and len(port) <= 5) or addr.endswith(":"):
        return None
    return addr.strip(), port


def _ltm_search(qs, search: str):
    """SLB 页搜索：VS 名 / 虚拟地址 / 池名 / 设备名 直接匹配；成员走反查拼进 OR。

    词里带端口（``vsaddress:port`` / ``poolmemberaddress:port``）时拆成
    地址与端口两个条件**同时**命中——跨字段组合是整串 icontains 做不到的：

    - VS 侧：``vs_address`` 命中地址 且 ``vs_port`` 精确等于端口；
    - 成员侧：成员表按 address + port 查出 (device, pool_name)，再折算回
      「该设备上引用这些池的 VS」（VS 行上没有成员字段，只能这样反查）。
    """
    q = (
        Q(name__icontains=search)
        | Q(vs_address__icontains=search)
        | Q(pool__icontains=search)
        | Q(device__hostname__icontains=search)
    )
    pair = _split_addr_port(search)
    addr, port = pair if pair else (search, None)
    if pair:
        # 地址:端口 组合——VS 自身命中，或成员按 地址+端口 命中后反查
        q |= Q(vs_address__icontains=addr, vs_port=port)
    member_q = Q(address__icontains=addr)
    if port:
        member_q &= Q(port=port)
    hits = set(LtmPoolMember.objects.filter(member_q).values_list("device_id", "pool_name")[:MAX_SEARCH_PAIRS])
    if hits:
        # (device_id, pool_name) → 该设备上 pool 名命中的 VS
        q |= _pairs_q(hits, value_field="pool")
    return qs.filter(q)


def _ltm_rows(page: list[LtmVirtualServer]) -> list[dict]:
    """按页批量取池与成员，拼成 VS 行（每类一次查询，不逐行查库）。"""
    pool_keys = {(vs.device_id, vs.pool) for vs in page if vs.pool}
    device_ids = {vs.device_id for vs in page}
    pool_names = {name for _, name in pool_keys}

    pools: dict[tuple[int, str], LtmPool] = {}
    if pool_keys:
        pools = {
            (p.device_id, p.name): p
            for p in LtmPool.objects.filter(device_id__in=device_ids, name__in=pool_names).select_related("device")
        }

    members: dict[tuple[int, str], list[LtmPoolMember]] = {}
    if pool_names:
        for m in LtmPoolMember.objects.filter(device_id__in=device_ids, pool_name__in=pool_names):
            members.setdefault((m.device_id, m.pool_name), []).append(m)

    rows = []
    for vs in page:
        pool = pools.get((vs.device_id, vs.pool)) if vs.pool else None
        rows.append(
            {
                "device": vs.device_id,
                "device_hostname": vs.device.hostname,
                "name": vs.name,
                "vs_address": vs.vs_address,
                "vs_port": vs.vs_port or "",
                "protocol": vs.protocol or "",
                "status": vs.status,
                "snat_type": vs.snat_type or "",
                "persist": vs.persist or "",
                # 一级 VS 行单独展示的字段：profile / iRule 是模型上的 JSON 名字列表
                "profiles": vs.profiles or [],
                "rules": vs.rules or [],
                "pool": (
                    {
                        "name": pool.name,
                        "mode": pool.mode,
                        "monitors": pool.monitors or [],
                    }
                    if pool
                    else None
                ),
                "members": [
                    {"name": m.name, "address": m.address, "port": m.port or ""}
                    for m in members.get((vs.device_id, vs.pool), [])
                ],
            }
        )
    return rows


@api_view(["GET"])
@permission_classes([AllowAny])
def ltm_chains(request):
    """负载均衡关联链：GET /api/lb-chain/slb/?device=&search=&ordering=&page="""
    qs = _device_filter(request, LtmVirtualServer.objects.select_related("device").all())
    search = (request.query_params.get("search") or "").strip()
    if search:
        qs = _ltm_search(qs, search)
    qs = _order(qs, request, allowed_extra={"device_hostname": "device__hostname"})

    page, paginator = _page(request, qs)
    if page is None:  # 页码非法时 paginate_queryset 已经抛 NotFound，这里只兜空页
        page = []
    return paginator.get_paginated_response(_ltm_rows(page))


# ---------------------------------------------------------------------------
# GTM：WideIP → 池 → GTM 虚拟服务器
# ---------------------------------------------------------------------------


def _wideip_pairs_q(pairs) -> Q:
    """(device_id, 池名) → wideip.pools JSON 包含该池名（逐对 OR，JSON 语义没法拆成 __in）。"""
    q = Q()
    for device_id, pool_name in pairs:
        q |= Q(device_id=device_id, pools__contains=[pool_name])
    return q


def _wideip_pairs(pool_q) -> set[tuple[int, str]]:
    """按「对 GtmPool 的过滤条件」反查出 (device_id, 池名) 集合。"""
    if not pool_q:
        return set()
    return set(GtmPool.objects.filter(pool_q).values_list("device_id", "name")[:MAX_SEARCH_PAIRS])


def _pool_q_via_servers(name_icontains: str = "", monitor_icontains: str = "") -> Q:
    """GtmServer 名 / 健康检查 命中 → 池条件（池成员以 server 名引用它）。

    ``members__contains`` 必须传**数组形式** ``[{...}]``：本项目实测这台 PG 的
    jsonb 数组 ``@> 对象`` 不生效（返回 0），``@> [对象]`` 才命中。
    """
    cond = Q()
    if name_icontains:
        cond |= Q(name__icontains=name_icontains)
    if monitor_icontains:
        cond |= Q(monitor__icontains=monitor_icontains)
    q = Q()
    if cond:
        servers = GtmServer.objects.filter(cond).values_list("device_id", "name")[:MAX_SEARCH_PAIRS]
        for device_id, server_name in servers:
            q |= Q(device_id=device_id, members__contains=[{"server": server_name}])
    return q


def _pool_q_via_vservers(vs_cond: Q) -> Q:
    """GtmVServer 命中（名字/地址/端口/健康检查）→ 池条件（成员以 server+vs 引用它）。"""
    q = Q()
    triples = GtmVServer.objects.filter(vs_cond).values_list("device_id", "server__name", "name")[:MAX_SEARCH_PAIRS]
    for device_id, server_name, vs_name in triples:
        q |= Q(device_id=device_id, members__contains=[{"server": server_name, "vserver": vs_name}])
    return q


def _gtm_search(qs, search: str):
    """域名解析页搜索，覆盖形态：

    wideipname / poolname / servername / vservername / vserveraddress /
    vserveraddress:port / 健康检查名（池、服务器、虚拟服务器三处的 monitor）。

    除域名与设备名直接匹配外，其余都是**沿链反查**：命中对象 → 它所在/被引用的
    池 → 引用这些池的 wideip（``_wideip_pairs_q`` 的 JSON 包含）。
    """
    pair = _split_addr_port(search)
    q = Q(name__icontains=search) | Q(device__hostname__icontains=search)

    # 对 GtmPool 的反查条件：池名 + 池健康检查 + server 名/健康检查 + vserver 名/地址/健康检查
    pool_q = Q(name__icontains=search) | Q(monitor__icontains=search)
    pool_q |= _pool_q_via_servers(name_icontains=search, monitor_icontains=search)
    vs_cond = Q(name__icontains=search) | Q(ip_address__icontains=search) | Q(monitor__icontains=search)
    if pair:
        # vserveraddress:port —— 地址模糊 + 端口精确，跨字段组合
        vs_cond |= Q(ip_address__icontains=pair[0], port=pair[1])
    pool_q |= _pool_q_via_vservers(vs_cond)

    hits = _wideip_pairs(pool_q)
    if hits:
        q |= _wideip_pairs_q(hits)
    return qs.filter(q)


def _gtm_rows(page: list[GtmWideip]) -> list[dict]:
    """按页批量取 GTM 池与虚拟服务器，把成员折算出 IP/端口。

    成员只存 ``server``/``vserver`` 名字，IP 在 ``GtmVServer`` 上，
    需要 (device, server 名, vs 名) 三元组才能定位——按页建索引一次查齐。
    """
    device_ids = {w.device_id for w in page}
    wideip_pool_names = {name for w in page for name in (w.pools or [])}

    pools: dict[tuple[int, str], GtmPool] = {}
    if wideip_pool_names:
        pools = {
            (p.device_id, p.name): p
            for p in GtmPool.objects.filter(device_id__in=device_ids, name__in=wideip_pool_names)
        }

    # (device_id, server 名, vs 名) → GtmVServer；server 是 FK，名字与数据中心在 GtmServer 上
    vserver_index: dict[tuple[int, str, str], GtmVServer] = {}
    server_dc: dict[tuple[int, str], str] = {}
    server_names: dict[int, str] = {}
    if device_ids:
        # 一次查询同时建 id→名（vserver 索引用）与 (device_id, 名)→数据中心（成员列用）两个索引；
        # 后者的键必须是 device_id——写成 server.id 会在两者恰好相等时「蒙对」，套跑即错
        servers_qs = GtmServer.objects.filter(device_id__in=device_ids).values_list(
            "id", "device_id", "name", "datacenter"
        )
        for sid, sdev, sname, sdc in servers_qs.iterator():
            server_names[sid] = sname
            server_dc[(sdev, sname)] = sdc or ""
        for vs in GtmVServer.objects.filter(device_id__in=device_ids).select_related("server"):
            vserver_index[(vs.device_id, server_names.get(vs.server_id) or vs.server.name, vs.name)] = vs

    rows = []
    for w in page:
        pool_out = []
        for pool_name in w.pools or []:
            pool = pools.get((w.device_id, pool_name))
            member_out = []
            if pool:
                for entry in pool.members or []:
                    server_name, vs_name, raw = _parse_member_entry(entry)
                    target = vserver_index.get((w.device_id, server_name, vs_name))
                    # order / ratio / monitor：入库形态是 Saver 归一后的键，
                    # 手工 payload / 旧模板给的是 member_* 前缀，两种都要认
                    member_out.append(
                        {
                            "server": server_name,
                            "vserver": vs_name,
                            "status": str(raw.get("status") or raw.get("member_status") or ""),
                            "order": raw.get("order", raw.get("member_order")),
                            "ratio": raw.get("ratio", raw.get("member_ratio")),
                            "monitor": str(raw.get("monitor") or raw.get("member_monitor") or ""),
                            # 找不到 GtmVServer 就没有 IP——前端回退显示 server/vserver 名字；
                            # 数据中心只挂在 server 上，vs 缺失时仍能给出
                            "address": target.ip_address if target else None,
                            "port": (target.port if target else "") or "",
                            "found": target is not None,
                            "datacenter": server_dc.get((w.device_id, server_name), ""),
                        }
                    )
            pool_out.append(
                {
                    "name": pool_name,
                    # wideip 引用了但池记录不存在：给空骨架，前端才能显示「池缺失」而不是吞掉
                    "lb_mode": pool.lb_mode if pool else "",
                    "alternate_mode": pool.alternate_mode if pool else "",
                    "fallback_mode": pool.fallback_mode if pool else "",
                    "fallback_ip": (pool.fallback_ip if pool else None) or "",
                    "ttl": pool.ttl if pool else None,
                    "monitor": (pool.monitor if pool else []) or [],
                    "members": member_out,
                }
            )
        rows.append(
            {
                "device": w.device_id,
                "device_hostname": w.device.hostname,
                "name": w.name,
                "rtype": w.rtype or "",
                "lb_mode": w.lb_mode or "",
                "pools": pool_out,
            }
        )
    return rows


def _gtm_filter_monitor(qs, monitor: str):
    """健康检查类型过滤：池 / 服务器 / 虚拟服务器三处 monitor 任一命中即可。"""
    pool_q = Q(monitor__icontains=monitor)
    pool_q |= _pool_q_via_servers(monitor_icontains=monitor)
    pool_q |= _pool_q_via_vservers(Q(monitor__icontains=monitor))
    hits = _wideip_pairs(pool_q)
    # 反查不到任何池 → 该健康检查类型下没有 wideip，给空集而不是不过滤
    return qs.filter(_wideip_pairs_q(hits)) if hits else qs.none()


@api_view(["GET"])
@permission_classes([AllowAny])
def gtm_chains(request):
    """域名解析关联链：
    GET /api/lb-chain/gslb/?device=&rtype=&monitor=&search=&ordering=&page=

    - ``rtype``：WideIP 记录类型精确过滤（下拉）；
    - ``monitor``：健康检查类型过滤，池/服务器/虚拟服务器三处 monitor 命中其一即可。
    """
    qs = _device_filter(request, GtmWideip.objects.select_related("device").all())
    rtype = (request.query_params.get("rtype") or "").strip()
    if rtype:
        qs = qs.filter(rtype__iexact=rtype)
    monitor = (request.query_params.get("monitor") or "").strip()
    if monitor:
        qs = _gtm_filter_monitor(qs, monitor)
    search = (request.query_params.get("search") or "").strip()
    if search:
        qs = _gtm_search(qs, search)
    qs = _order(qs, request, allowed_extra={"device_hostname": "device__hostname"})

    page, paginator = _page(request, qs)
    if page is None:
        page = []
    return paginator.get_paginated_response(_gtm_rows(page))


@api_view(["GET"])
@permission_classes([AllowAny])
def gtm_facets(request):
    """过滤下拉选项：GET /api/lb-chain/gslb/facets/

    ``rtypes`` 来自 wideip 记录类型去重；``monitors`` 是健康检查类型合集——
    池（JSON 列表，Python 侧展开）+ 服务器 / 虚拟服务器（CharField）三处 distinct。
    """
    rtypes = sorted({v for v in GtmWideip.objects.values_list("rtype", flat=True) if v})
    monitors: set[str] = set()
    for lst in GtmPool.objects.values_list("monitor", flat=True).iterator():
        monitors.update(m for m in (lst or []) if isinstance(m, str) and m)
    monitors.update(v for v in GtmServer.objects.values_list("monitor", flat=True) if v)
    monitors.update(v for v in GtmVServer.objects.values_list("monitor", flat=True) if v)
    return Response({"rtypes": rtypes, "monitors": sorted(monitors, key=str.lower)})


# ---------------------------------------------------------------------------
# 导出：扁平宽表 xlsx（openpyxl 后端生成、前端只下载 Blob——与互联网资产分析同模式）
#
# **列序与前端 slb.vue / gslb.vue 的扁平列组手工保持一致**（表格渲染在前端、
# 导出在后端的两份扁平化并存是项目既定先例，见 internet-asset 的 build_path_rows）。
# ---------------------------------------------------------------------------

SLB_EXPORT_HEADERS = [
    "名称",  # 成员地址#端口；回退行是池名 / VS地址#端口
    "设备",
    "VS地址#端口",
    "VS名称",
    "协议",
    "SNAT",
    "会话保持",
    "Profile",
    "iRule",
    "关联池",
    "池负载模式",
    "池监控",
]
SLB_EXPORT_COLUMN_WIDTHS = (24, 16, 22, 24, 8, 12, 12, 30, 30, 16, 14, 22)

GSLB_EXPORT_HEADERS = [
    "名称",  # 成员地址#端口；回退行是池名 / 域名
    "设备",
    "域名",
    "记录类型",
    "WideIP算法",
    "池名",
    "池算法",
    "fallback",
    "TTL",
    "池监控",
    "池Order",
    "池Ratio",
    "成员Order",
    "成员Ratio",
    "成员监控",
    "数据中心",
    "状态",
]
GSLB_EXPORT_COLUMN_WIDTHS = (24, 16, 24, 10, 14, 16, 24, 20, 8, 22, 10, 10, 12, 12, 14, 14, 10)

_STATE_LABELS = {"ok": "正常", "disabled": "停用", "lost": "未找到"}


def _ltm_flat_rows(chain_rows: list[dict]) -> list[list]:
    """SLB 扁平宽表：以链最深层为行（成员→池→VS 回退），VS/池字段整条下填。"""
    rows: list[list] = []
    for r in chain_rows:
        vs_label = _join_ip_port(r["vs_address"], r["vs_port"]) if r["vs_address"] else r["name"]
        base = [
            r["device_hostname"],
            vs_label,
            r["name"],
            r["protocol"] or "-",
            r["snat_type"] or "-",
            r["persist"] or "-",
            "、".join(r["profiles"]) or "-",
            "、".join(r["rules"]) or "-",
        ]
        pool = r.get("pool")
        pool_seg = [pool["name"], pool["mode"] or "-", "、".join(pool["monitors"]) or "-"] if pool else ["-", "-", "-"]
        if pool and r["members"]:
            for m in r["members"]:
                label = _join_ip_port(m["address"], m["port"]) if m["address"] else m["name"]
                rows.append([label] + base + pool_seg)
        else:
            rows.append([pool["name"] if pool else vs_label] + base + pool_seg)
    return rows


def _gtm_flat_rows(chain_rows: list[dict]) -> list[list]:
    """GSLB 扁平宽表：粒度规则同 LTM；池级 Order/Ratio = 池内成员取值集合去重。"""
    rows: list[list] = []
    for w in chain_rows:
        wide = [w["device_hostname"], w["name"], w["rtype"] or "-", w["lb_mode"] or "-"]
        empty_pool = ["-"] * 7  # 池名/池算法/fallback/TTL/池监控/池Order/池Ratio
        empty_member = ["", "", "", "", "-"]  # 成员Order/成员Ratio/成员监控/数据中心/状态
        if not w["pools"]:
            rows.append([w["name"]] + wide + empty_pool + empty_member)
            continue
        for p in w["pools"]:
            algo = " / ".join(x for x in (p["lb_mode"], p["alternate_mode"]) if x) or "-"
            fallback = (p["fallback_mode"] or "") + (f"（{p['fallback_ip']}）" if p["fallback_ip"] else "")
            pool_seg = [
                p["name"],
                algo,
                fallback or "-",
                str(p["ttl"]) if p["ttl"] is not None else "-",
                "、".join(p["monitor"]) or "-",
                "、".join(str(x) for x in sorted({m["order"] for m in p["members"] if m["order"] is not None})) or "-",
                "、".join(str(x) for x in sorted({m["ratio"] for m in p["members"] if m["ratio"] is not None})) or "-",
            ]
            if not p["members"]:
                rows.append([p["name"]] + wide + pool_seg + empty_member)
                continue
            for m in p["members"]:
                label = (
                    _join_ip_port(m["address"], m["port"])
                    if m["found"] and m["address"]
                    else f"{m['server']}/{m['vserver']}"
                )
                state_key = "lost" if not m["found"] else ("disabled" if m["status"] == "disabled" else "ok")
                member_seg = [
                    "" if m["order"] is None else str(m["order"]),
                    "" if m["ratio"] is None else str(m["ratio"]),
                    m["monitor"],
                    m["datacenter"],
                    _STATE_LABELS.get(state_key, "-"),
                ]
                rows.append([label] + wide + pool_seg + member_seg)
    return rows


def _xlsx_response(headers: list[str], widths: tuple, rows: list[list], title: str, filename: str):
    """表头 + 行写进 xlsx 返回附件响应（列宽按 get_column_letter 逐列设置）。"""
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        sheet = workbook.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    # 列宽跟着表头走：写死列字母在加列时会静默少设宽度（与互联网资产分析同约定）
    assert len(widths) == len(headers)  # noqa: S101 —— 测试里另有断言守，这里尽早炸
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


@api_view(["GET"])
@permission_classes([AllowAny])
def ltm_export(request):
    """导出 SLB 扁平宽表：GET /api/lb-chain/slb/export/?device=&search=

    全量导出（不分页），过滤/搜索参数与列表接口同语义。
    """
    qs = _device_filter(request, LtmVirtualServer.objects.select_related("device").all())
    search = (request.query_params.get("search") or "").strip()
    if search:
        qs = _ltm_search(qs, search)
    rows = _ltm_flat_rows(_ltm_rows(list(qs.order_by(*DEFAULT_ORDER))))
    return _xlsx_response(
        SLB_EXPORT_HEADERS,
        SLB_EXPORT_COLUMN_WIDTHS,
        rows,
        "负载均衡关联链",
        f"lb-chains-{date.today():%Y%m%d}.xlsx",
    )


@api_view(["GET"])
@permission_classes([AllowAny])
def gtm_export(request):
    """导出 GSLB 扁平宽表：GET /api/lb-chain/gslb/export/?device=&rtype=&monitor=&search="""
    qs = _device_filter(request, GtmWideip.objects.select_related("device").all())
    rtype = (request.query_params.get("rtype") or "").strip()
    if rtype:
        qs = qs.filter(rtype__iexact=rtype)
    monitor = (request.query_params.get("monitor") or "").strip()
    if monitor:
        qs = _gtm_filter_monitor(qs, monitor)
    search = (request.query_params.get("search") or "").strip()
    if search:
        qs = _gtm_search(qs, search)
    rows = _gtm_flat_rows(_gtm_rows(list(qs.order_by(*DEFAULT_ORDER))))
    return _xlsx_response(
        GSLB_EXPORT_HEADERS,
        GSLB_EXPORT_COLUMN_WIDTHS,
        rows,
        "域名解析关联链",
        f"dn-chains-{date.today():%Y%m%d}.xlsx",
    )
