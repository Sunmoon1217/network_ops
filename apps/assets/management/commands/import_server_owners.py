"""从 xlsx 导入「服务器 IP ↔ 负责人」。

    python manage.py import_server_owners --file servers.xlsx
    python manage.py import_server_owners --file servers.xlsx --dry-run
    python manage.py import_server_owners --file servers.xlsx --sheet servers

表结构：sheet 名 ``servers``，表头列 ``hostname`` / ``ip`` / ``owner``。

几点约定：

- **按列名取，不按位置**：表头顺序可以变，也能容忍多出来的列。缺列直接报错并
  列出实际表头，而不是静默按位置串位（``ip`` 串到 ``hostname`` 上很难发现）。
- **``ip`` 是唯一键**：同一个 IP 再次导入就是更新；文件内重复取最后一行。
- **不动 ``status``**：这份表里没有这一列，所以更新时保持原值，避免把手工停用的
  记录又导成启用。
- 有非法行（``ip`` 为空或不是合法地址）时**整批不导入**并以非零退出码收尾。
  不做「写一半再报错」：那样操作者不知道到底落了哪些行，比全不导更难处理。
  整行空白只是跳过，不算错误。
"""

import ipaddress
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

SHEET_NAME = "servers"
REQUIRED_COLUMNS = ("hostname", "ip", "owner")


class Command(BaseCommand):
    help = "从 xlsx 导入服务器负责人（sheet: servers，列: hostname / ip / owner）"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, metavar="PATH", help="xlsx 文件路径")
        parser.add_argument("--sheet", default=SHEET_NAME, metavar="NAME", help=f"工作表名，默认 {SHEET_NAME}")
        parser.add_argument("--dry-run", action="store_true", help="只校验并试算增改数，不写库")

    def handle(self, *args, **options):
        from openpyxl import load_workbook

        from assets.models import ServerOwner

        path = Path(options["file"])
        if not path.is_file():
            raise CommandError(f"文件不存在: {path}")

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_name = options["sheet"]
            if sheet_name not in workbook.sheetnames:
                raise CommandError(f"工作表 {sheet_name!r} 不存在，文件里有: {', '.join(workbook.sheetnames)}")
            valid, errors, duplicates = self._read(workbook[sheet_name])
        finally:
            workbook.close()

        if duplicates:
            self.stdout.write(self.style.WARNING(f"文件内重复的 ip {len(duplicates)} 个，取最后一行: {duplicates}"))
        for message in errors:
            self.stdout.write(self.style.ERROR(f"  {message}"))

        if errors:
            # 先校验后写：有一行不合法就整批不动，避免留下「导了一半」的状态
            self.stdout.write("")
            raise CommandError(f"{len(errors)} 行数据有问题，整批未导入（见上面的行号）")

        to_create, to_update = self._split(ServerOwner, valid)
        if options["dry_run"]:
            self.stdout.write(self.style.WARNING("--dry-run：没有写库"))
        else:
            self._apply(ServerOwner, to_create, to_update)

        self.stdout.write("")
        verb = "将新建" if options["dry_run"] else "新建"
        self.stdout.write(
            f"完成：{verb} {len(to_create)} 条，"
            f"{'将更新' if options['dry_run'] else '更新'} {len(to_update)} 条，"
            f"文件内重复 {len(duplicates)} 个 ip"
        )

    def _read(self, sheet):
        """按列名解析，返回 ``(有效行, 错误清单, 重复的 ip)``。

        有效行是 ``(xlsx 行号, hostname, ip, owner)``，``ip`` 已规范化成压缩小写形式。
        """
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise CommandError("工作表是空的")

        index = {}
        for position, value in enumerate(header):
            name = str(value or "").strip().lower()
            if name and name not in index:
                index[name] = position

        missing = [column for column in REQUIRED_COLUMNS if column not in index]
        if missing:
            raise CommandError(f"缺少列: {', '.join(missing)}；表头实际是: {', '.join(sorted(index)) or '(空)'}")

        by_ip: dict[str, tuple] = {}
        duplicates: list[str] = []
        errors: list[str] = []

        for row_no, row in enumerate(rows, start=2):
            hostname = self._cell(row, index, "hostname")
            raw_ip = self._cell(row, index, "ip")
            owner = self._cell(row, index, "owner")

            if not any((hostname, raw_ip, owner)):
                continue  # 整行空白

            if not raw_ip:
                errors.append(f"第 {row_no} 行: ip 为空")
                continue
            try:
                # 规范化成压缩小写，避免 2001:DB8::1 与 2001:db8::1 落成两条
                ip = str(ipaddress.ip_address(raw_ip))
            except ValueError:
                errors.append(f"第 {row_no} 行: ip 不是合法地址 {raw_ip!r}")
                continue

            if ip in by_ip:
                duplicates.append(ip)
            by_ip[ip] = (row_no, hostname, ip, owner)

        return list(by_ip.values()), errors, sorted(set(duplicates))

    def _cell(self, row, index: dict[str, int], column: str) -> str:
        position = index[column]
        if position >= len(row):
            return ""
        value = row[position]
        return str(value).strip() if value is not None else ""

    def _split(self, model, valid: list[tuple]) -> tuple[list, list]:
        """按 ip 分成待新建 / 待更新，更新只覆盖 hostname 与 owner（不动 status）"""
        existing = {item.ip: item for item in model.objects.all()}

        to_create, to_update = [], []
        now = timezone.now()
        for _, hostname, ip, owner in valid:
            item = existing.get(ip)
            if item is None:
                to_create.append(model(hostname=hostname, ip=ip, owner=owner))
            else:
                item.hostname = hostname
                item.owner = owner
                item.updated_at = now
                to_update.append(item)
        return to_create, to_update

    def _apply(self, model, to_create: list, to_update: list) -> None:
        with transaction.atomic():
            if to_create:
                model.objects.bulk_create(to_create, batch_size=500)
            if to_update:
                # bulk_update 不走 pre_save，auto_now 的 updated_at 上面已经手动填了
                model.objects.bulk_update(to_update, ["hostname", "owner", "updated_at"], batch_size=500)
