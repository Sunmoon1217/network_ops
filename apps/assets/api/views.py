import logging
from pathlib import Path

from django.http import JsonResponse
from rest_framework import viewsets
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from assets.models import (
    ArpMac,
    AddressBook,
    Cabinet,
    DataCenter,
    Device,
    DeviceAccount,
    DeviceConfig,
    DeviceConnection,
    DeviceModel,
    GtmDatacenter,
    GtmPool,
    GtmWideip,
    Interface,
    LtmPool,
    LtmPoolMember,
    LtmVirtualServer,
    NatRule,
    NtpConfig,
    Policy,
    Room,
    Route,
    SecurityZone,
    Service,
    SnmpConfig,
    Subnet,
    SyslogConfig,
    Vendor,
    Vlan,
    Vrf,
)
from ops.config_repo import save_config

from .serializers import (
    CabinetSerializer,
    DataCenterSerializer,
    DeviceAccountSerializer,
    DeviceConfigSerializer,
    DeviceConnectionSerializer,
    DeviceModelSerializer,
    DeviceSerializer,
    InterfaceSerializer,
    NtpConfigSerializer,
    RoomSerializer,
    SecurityZoneSerializer,
    SnmpConfigSerializer,
    SyslogConfigSerializer,
    VendorSerializer,
    VlanSerializer,
    VrfSerializer,
)

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Overview
# ---------------------------------------------------------------------------


def overview(request):
    """资产管理概览统计"""
    from django.db.models import Count

    # --- 基础设施 ---
    cabinets_by_dc = list(
        Cabinet.objects.values("room__datacenter__name")
        .annotate(count=Count("id"))
        .order_by("-count")[:10]
    )
    cabinets_by_status = list(
        Cabinet.objects.values("status").annotate(count=Count("id"))
    )

    # --- 设备 ---
    device_types = list(
        Device.objects.values("device_type").annotate(count=Count("id")).order_by("-count")
    )
    devices_by_vendor = list(
        Device.objects.values("device_model__vendor__name")
        .annotate(count=Count("id")).order_by("-count")[:10]
    )
    devices_by_zone = list(
        Device.objects.values("security_zone__name").annotate(count=Count("id")).order_by("-count")
    )

    # --- 接口 ---
    interface_modes = list(
        Interface.objects.values("mode").annotate(count=Count("id"))
    )

    # --- 路由 ---
    route_protocols = list(
        Route.objects.values("protocol").annotate(count=Count("id")).order_by("-count")
    )
    routes_by_device = list(
        Route.objects.values("vrf__device__hostname")
        .annotate(count=Count("id")).order_by("-count")[:10]
    )

    # --- 防火墙 ---
    policies_by_action = list(
        Policy.objects.values("action").annotate(count=Count("id"))
    )
    nat_by_type = list(
        NatRule.objects.values("nat_type").annotate(count=Count("id"))
    )
    address_books_by_type = list(
        AddressBook.objects.values("address_type").annotate(count=Count("id"))
    )
    services_by_protocol = list(
        Service.objects.values("protocol").annotate(count=Count("id"))
    )

    # --- IPAM ---
    subnets_by_dc = list(
        Subnet.objects.values("datacenter__name").annotate(count=Count("id")).order_by("-count")
    )
    subnets_by_zone = list(
        Subnet.objects.values("security_zone__name").annotate(count=Count("id")).order_by("-count")
    )

    # --- SLB ---
    vs_by_protocol = list(
        LtmVirtualServer.objects.values("protocol").annotate(count=Count("id"))
    )
    pools_by_lb = list(
        LtmPool.objects.values("device__hostname").annotate(count=Count("id")).order_by("-count")[:10]
    )

    return JsonResponse({
        # 基础设施
        "dc_count": DataCenter.objects.count(),
        "room_count": Room.objects.count(),
        "cabinet_count": Cabinet.objects.count(),
        "cabinet_active": Cabinet.objects.filter(status="active").count(),
        "cabinets_by_dc": cabinets_by_dc,

        "cabinets_by_status": cabinets_by_status,
        "security_zone_count": SecurityZone.objects.count(),

        # 设备
        "device_count": Device.objects.count(),
        "device_types": device_types,
        "vendor_count": Vendor.objects.count(),
        "device_model_count": DeviceModel.objects.count(),
        "devices_by_vendor": devices_by_vendor,
        "devices_by_zone": devices_by_zone,

        # 接口
        "interface_count": Interface.objects.count(),
        "interface_up": Interface.objects.filter(enabled=True).count(),
        "interface_modes": interface_modes,

        # 网络
        "vlan_count": Vlan.objects.count(),
        "vrf_count": Vrf.objects.count(),
        "route_count": Route.objects.count(),
        "route_protocols": route_protocols,
        "routes_by_device": routes_by_device,

        # 负载均衡
        "vs_count": LtmVirtualServer.objects.count(),
        "pool_count": LtmPool.objects.count(),
        "pool_member_count": LtmPoolMember.objects.count(),
        "vs_by_protocol": vs_by_protocol,
        "pools_by_lb": pools_by_lb,

        # 全局负载均衡
        "gtm_dc_count": GtmDatacenter.objects.count(),
        "gtm_wideip_count": GtmWideip.objects.count(),
        "gtm_pool_count": GtmPool.objects.count(),

        # 防火墙
        "address_book_count": AddressBook.objects.count(),
        "service_count": Service.objects.count(),
        "policy_count": Policy.objects.count(),
        "nat_rule_count": NatRule.objects.count(),
        "policies_by_action": policies_by_action,
        "nat_by_type": nat_by_type,
        "address_books_by_type": address_books_by_type,
        "services_by_protocol": services_by_protocol,

        # IPAM
        "subnet_count": Subnet.objects.count(),
        "arp_mac_count": ArpMac.objects.count(),
        "subnets_by_dc": subnets_by_dc,
        "subnets_by_zone": subnets_by_zone,
    })


# ---------------------------------------------------------------------------
# DCIM ViewSets
# ---------------------------------------------------------------------------


class SecurityZoneViewSet(viewsets.ModelViewSet):
    queryset = SecurityZone.objects.all()
    serializer_class = SecurityZoneSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name",)
    ordering_fields = ("name", "created_at")


class DataCenterViewSet(viewsets.ModelViewSet):
    queryset = DataCenter.objects.all()
    serializer_class = DataCenterSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "address")
    ordering_fields = ("name", "created_at")


class RoomViewSet(viewsets.ModelViewSet):
    queryset = Room.objects.select_related("datacenter").all()
    serializer_class = RoomSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "datacenter__name")
    ordering_fields = ("name", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        datacenter = self.request.query_params.get("datacenter")
        if datacenter:
            qs = qs.filter(datacenter_id=datacenter)
        return qs


class CabinetViewSet(viewsets.ModelViewSet):
    queryset = Cabinet.objects.select_related("room", "room__datacenter").all()
    serializer_class = CabinetSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "room__name", "room__datacenter__name")
    ordering_fields = ("name", "status", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        room = self.request.query_params.get("room")
        if room:
            qs = qs.filter(room_id=room)
        datacenter = self.request.query_params.get("datacenter")
        if datacenter:
            qs = qs.filter(room__datacenter_id=datacenter)
        status = self.request.query_params.get("status")
        if status:
            qs = qs.filter(status=status)
        return qs


class VendorViewSet(viewsets.ModelViewSet):
    queryset = Vendor.objects.all()
    serializer_class = VendorSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name",)
    ordering_fields = ("name",)


class DeviceModelViewSet(viewsets.ModelViewSet):
    queryset = DeviceModel.objects.select_related("vendor").all()
    serializer_class = DeviceModelSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "vendor__name")
    ordering_fields = ("name",)


class DeviceViewSet(viewsets.ModelViewSet):
    queryset = Device.objects.select_related("idc", "cabinet", "security_zone", "device_model").all()
    serializer_class = DeviceSerializer
    permission_classes = (AllowAny,)
    search_fields = ("hostname", "ip_address")
    ordering_fields = ("hostname", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_type = self.request.query_params.get("device_type")
        if device_type:
            qs = qs.filter(device_type=device_type)
        idc = self.request.query_params.get("idc")
        if idc:
            qs = qs.filter(idc_id=idc)
        return qs


class DeviceConfigViewSet(viewsets.ModelViewSet):
    queryset = DeviceConfig.objects.select_related("device").all()
    serializer_class = DeviceConfigSerializer
    permission_classes = (AllowAny,)
    ordering_fields = ("collected_at",)

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class DeviceConnectionViewSet(viewsets.ModelViewSet):
    queryset = DeviceConnection.objects.select_related("device").all()
    serializer_class = DeviceConnectionSerializer
    permission_classes = (AllowAny,)
    search_fields = ("username", "device__hostname")
    ordering_fields = ("created_at",)

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


# ---------------------------------------------------------------------------
# Network ViewSets
# ---------------------------------------------------------------------------


class VlanViewSet(viewsets.ModelViewSet):
    queryset = Vlan.objects.all()
    serializer_class = VlanSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "vid")
    ordering_fields = ("vid", "name")


class VrfViewSet(viewsets.ModelViewSet):
    queryset = Vrf.objects.select_related("device").all()
    serializer_class = VrfSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")
    ordering_fields = ("name", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class InterfaceViewSet(viewsets.ModelViewSet):
    queryset = Interface.objects.select_related("device", "vrf").all()
    serializer_class = InterfaceSerializer
    permission_classes = (AllowAny,)
    search_fields = ("interface", "device__hostname", "ip_address")
    ordering_fields = ("interface", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        mode = self.request.query_params.get("mode")
        if mode:
            qs = qs.filter(mode=mode)
        return qs


class DeviceAccountViewSet(viewsets.ModelViewSet):
    queryset = DeviceAccount.objects.select_related("device").all()
    serializer_class = DeviceAccountSerializer
    permission_classes = (AllowAny,)
    search_fields = ("username", "device__hostname")
    ordering_fields = ("username", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


# ---------------------------------------------------------------------------
# Baseline ViewSets
# ---------------------------------------------------------------------------


class SnmpConfigViewSet(viewsets.ModelViewSet):
    queryset = SnmpConfig.objects.select_related("device").all()
    serializer_class = SnmpConfigSerializer
    permission_classes = (AllowAny,)
    ordering_fields = ("created_at",)


class NtpConfigViewSet(viewsets.ModelViewSet):
    queryset = NtpConfig.objects.select_related("device").all()
    serializer_class = NtpConfigSerializer
    permission_classes = (AllowAny,)
    ordering_fields = ("created_at",)


class SyslogConfigViewSet(viewsets.ModelViewSet):
    queryset = SyslogConfig.objects.select_related("device").all()
    serializer_class = SyslogConfigSerializer
    permission_classes = (AllowAny,)
    ordering_fields = ("created_at",)


# ---------------------------------------------------------------------------
# Connection config map
# ---------------------------------------------------------------------------

CONNECTION_CONFIG_MAP = {
    ("H3C", "switch"): {"connection_type": "napalm", "driver": "h3c_comware"},
    ("H3C", "router"): {"connection_type": "napalm", "driver": "h3c_comware"},
    ("锐捷", "switch"): {"connection_type": "netmiko", "driver": "ruijie_os"},
    ("思科", "firewall"): {"connection_type": "napalm", "driver": "asa"},
    ("山石", "firewall"): {"connection_type": "netmiko", "driver": "hillstone_stoneos"},
    ("F5", "loadbalancer"): {"connection_type": "netmiko", "driver": "f5_ltm"},
    ("A10", "loadbalancer"): {"connection_type": "netmiko", "driver": "a10"},
}


def _get_connection_config(vendor_name, device_type):
    return CONNECTION_CONFIG_MAP.get(
        (vendor_name, device_type),
        {"connection_type": "netmiko", "driver": ""},
    )


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------


@api_view(["POST"])
@permission_classes([AllowAny])
def import_excel(request):
    """导入 Excel 文件，支持多个 Sheet"""
    from openpyxl import load_workbook

    excel_file = request.FILES.get("file")
    if not excel_file:
        return Response({"error": "请上传 Excel 文件"}, status=400)

    try:
        wb = load_workbook(excel_file, read_only=True, data_only=True)
    except Exception as e:
        return Response({"error": f"文件格式错误: {e}"}, status=400)

    results = {}
    sheet_importers = {
        "数据中心": _import_datacenters,
        "机房": _import_rooms,
        "机柜": _import_cabinets,
        "安全区": _import_security_zones,
        "设备": _import_devices,
        "配置文件": _import_configs,
    }

    for sheet_name, importer in sheet_importers.items():
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            try:
                results[sheet_name] = importer(ws)
            except Exception as e:
                logger.error("导入 %s 失败: %s", sheet_name, e)
                results[sheet_name] = {"created": 0, "updated": 0, "errors": [str(e)]}
        else:
            results[sheet_name] = {"created": 0, "updated": 0, "errors": [], "skipped": True}

    wb.close()
    return Response({"success": True, "results": results})


def _import_datacenters(ws):
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        name = str(row[0]).strip()
        defaults = {
            "address": str(row[1] or "").strip(),
            "contact": str(row[2] or "").strip(),
            "phone": str(row[3] or "").strip(),
            "remark": str(row[4] or "").strip(),
        }
        try:
            _, is_created = DataCenter.objects.update_or_create(name=name, defaults=defaults)
            created += 1 if is_created else 0
            updated += 0 if is_created else 1
        except Exception as e:
            errors.append(f"行 {row_idx}: {e}")
    return {"created": created, "updated": updated, "errors": errors}


def _import_rooms(ws):
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[1]:
            continue
        try:
            dc = DataCenter.objects.get(name=str(row[0]).strip())
            _, is_created = Room.objects.update_or_create(
                datacenter=dc,
                name=str(row[1]).strip(),
                defaults={"contact": str(row[2] or "").strip(), "remark": str(row[3] or "").strip()},
            )
            created += 1 if is_created else 0
            updated += 0 if is_created else 1
        except DataCenter.DoesNotExist:
            errors.append(f"行 {row_idx}: 数据中心 '{row[0]}' 不存在")
        except Exception as e:
            errors.append(f"行 {row_idx}: {e}")
    return {"created": created, "updated": updated, "errors": errors}


def _import_cabinets(ws):
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0] or not row[1] or not row[2]:
            continue
        try:
            dc = DataCenter.objects.get(name=str(row[0]).strip())
            room = Room.objects.get(datacenter=dc, name=str(row[1]).strip())
            _, is_created = Cabinet.objects.update_or_create(
                room=room,
                name=str(row[2]).strip(),
                defaults={
                    "row": str(row[3] or "").strip(),
                    "total_u": int(row[4] or 42),
                    "power_capacity": float(row[5]) if row[5] else None,
                    "status": str(row[6] or "active").strip(),
                    "remark": str(row[7] or "").strip(),
                },
            )
            created += 1 if is_created else 0
            updated += 0 if is_created else 1
        except DataCenter.DoesNotExist:
            errors.append(f"行 {row_idx}: 数据中心 '{row[0]}' 不存在")
        except Room.DoesNotExist:
            errors.append(f"行 {row_idx}: 机房 '{row[1]}' 不存在")
        except Exception as e:
            errors.append(f"行 {row_idx}: {e}")
    return {"created": created, "updated": updated, "errors": errors}


def _import_security_zones(ws):
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        try:
            _, is_created = SecurityZone.objects.update_or_create(
                name=str(row[0]).strip(),
                defaults={"color": str(row[1] or "#3b82f6").strip(), "description": str(row[2] or "").strip()},
            )
            created += 1 if is_created else 0
            updated += 0 if is_created else 1
        except Exception as e:
            errors.append(f"行 {row_idx}: {e}")
    return {"created": created, "updated": updated, "errors": errors}


def _import_devices(ws):
    created, updated, errors = 0, 0, []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue
        hostname = str(row[0]).strip()
        device_type = str(row[2] or "switch").strip()
        ip_address = str(row[1] or "").strip()

        vendor = None
        device_model = None
        vendor_name = str(row[3] or "").strip()
        model_name = str(row[4] or "").strip()
        if vendor_name:
            vendor, _ = Vendor.objects.get_or_create(name=vendor_name)
        if model_name and vendor:
            device_model, _ = DeviceModel.objects.get_or_create(name=model_name, vendor=vendor)

        dc = None
        cabinet = None
        dc_name = str(row[5] or "").strip()
        room_name = str(row[6] or "").strip()
        cab_name = str(row[7] or "").strip()
        if dc_name and room_name and cab_name:
            try:
                dc = DataCenter.objects.get(name=dc_name)
                room = Room.objects.get(datacenter=dc, name=room_name)
                cabinet = Cabinet.objects.get(room=room, name=cab_name)
            except (DataCenter.DoesNotExist, Room.DoesNotExist, Cabinet.DoesNotExist):
                errors.append(f"行 {row_idx}: 机柜路径 '{dc_name}/{room_name}/{cab_name}' 不存在")

        security_zone = None
        zone_name = str(row[8] or "").strip()
        if zone_name:
            security_zone = SecurityZone.objects.filter(name=zone_name).first()

        device_defaults = {
            "device_type": device_type,
            "ip_address": ip_address,
            "device_model": device_model,
            "idc": dc,
            "cabinet": cabinet,
            "security_zone": security_zone,
            "u_position": int(row[9]) if row[9] else None,
            "height": int(row[10] or 1),
            "remark": str(row[11] or "").strip(),
        }

        try:
            device, is_created = Device.objects.update_or_create(hostname=hostname, defaults=device_defaults)
            created += 1 if is_created else 0
            updated += 0 if is_created else 1
        except Exception as e:
            errors.append(f"行 {row_idx}: {e}")
            continue

        account_type = str(row[12] or "").strip()
        username = str(row[13] or "").strip()
        password = str(row[14] or "").strip()
        if username and password:
            conn_config = _get_connection_config(vendor_name, device_type)
            conn_defaults = {
                "connection_type": conn_config["connection_type"],
                "driver": conn_config["driver"],
                "account_type": account_type or "admin",
                "username": username,
                "password": password,
                "enable_password": str(row[15] or "").strip(),
                "port": int(row[16] or 22) if row[16] else 22,
                "timeout": int(row[17] or 30) if row[17] else 30,
            }
            try:
                DeviceConnection.objects.update_or_create(
                    device=device, account_type=account_type or "admin", defaults=conn_defaults
                )
            except Exception as e:
                errors.append(f"行 {row_idx} 连接: {e}")

    return {"created": created, "updated": updated, "errors": errors}


def _import_configs(ws):
    created, skipped, errors = 0, 0, []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row[0]:
            continue

        hostname = str(row[0]).strip()
        filename = row[1] and str(row[1]).strip() or hostname + ".txt"
        config_dir = str(row[2] or "").strip()

        try:
            device = Device.objects.get(hostname=hostname)
        except Device.DoesNotExist:
            errors.append(f"行 {row_idx}: 设备 '{hostname}' 不存在，请先导入设备")
            continue

        from ops.config_repo import CONFIG_REPO_PATH

        if config_dir:
            config_path = Path(config_dir) / filename
        else:
            config_path = CONFIG_REPO_PATH / hostname / filename

        if not config_path.exists():
            errors.append(f"行 {row_idx}: 配置文件不存在: {config_path}")
            continue

        try:
            config_text = config_path.read_text(encoding="utf-8")
        except Exception as e:
            errors.append(f"行 {row_idx}: 读取失败: {e}")
            continue

        if not config_text.strip():
            errors.append(f"行 {row_idx}: 配置文件为空")
            continue

        try:
            commit_hash = save_config(hostname, config_text, message=f"import: {hostname} config")
        except Exception as e:
            errors.append(f"行 {row_idx}: 保存到 Git 失败: {e}")
            continue

        DeviceConfig.objects.update_or_create(
            device=device,
            git_commit_hash=commit_hash,
            defaults={"config_json": {}},
        )
        created += 1

    return {"created": created, "skipped": skipped, "errors": errors}


# ---------------------------------------------------------------------------
# SLB (LTM) ViewSets
# ---------------------------------------------------------------------------


class LtmVirtualServerViewSet(viewsets.ModelViewSet):
    from assets.models import LtmVirtualServer

    from .serializers import LtmVirtualServerSerializer

    queryset = LtmVirtualServer.objects.select_related("device").all()
    serializer_class = LtmVirtualServerSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")
    ordering_fields = ("name", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class LtmPoolViewSet(viewsets.ModelViewSet):
    from assets.models import LtmPool

    from .serializers import LtmPoolSerializer

    queryset = LtmPool.objects.select_related("device").all()
    serializer_class = LtmPoolSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")
    ordering_fields = ("name", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class LtmPoolMemberViewSet(viewsets.ModelViewSet):
    from assets.models import LtmPoolMember

    from .serializers import LtmPoolMemberSerializer

    queryset = LtmPoolMember.objects.all()
    serializer_class = LtmPoolMemberSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "pool_name")


class LtmProfileViewSet(viewsets.ModelViewSet):
    from assets.models import LtmProfile

    from .serializers import LtmProfileSerializer

    queryset = LtmProfile.objects.select_related("device").all()
    serializer_class = LtmProfileSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class LtmIRuleViewSet(viewsets.ModelViewSet):
    from assets.models import LtmIRule

    from .serializers import LtmIRuleSerializer

    queryset = LtmIRule.objects.select_related("device").all()
    serializer_class = LtmIRuleSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class LtmSNATViewSet(viewsets.ModelViewSet):
    from assets.models import LtmSNAT

    from .serializers import LtmSNATSerializer

    queryset = LtmSNAT.objects.select_related("device").all()
    serializer_class = LtmSNATSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class LtmPersistViewSet(viewsets.ModelViewSet):
    from assets.models import LtmPersist

    from .serializers import LtmPersistSerializer

    queryset = LtmPersist.objects.select_related("device").all()
    serializer_class = LtmPersistSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


# ---------------------------------------------------------------------------
# GSLB (GTM) ViewSets
# ---------------------------------------------------------------------------


class GtmDatacenterViewSet(viewsets.ModelViewSet):
    from assets.models import GtmDatacenter

    from .serializers import GtmDatacenterSerializer

    queryset = GtmDatacenter.objects.select_related("device").all()
    serializer_class = GtmDatacenterSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class GtmWideipViewSet(viewsets.ModelViewSet):
    from assets.models import GtmWideip

    from .serializers import GtmWideipSerializer

    queryset = GtmWideip.objects.select_related("device").all()
    serializer_class = GtmWideipSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")
    ordering_fields = ("name", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class GtmPoolViewSet(viewsets.ModelViewSet):
    from assets.models import GtmPool

    from .serializers import GtmPoolSerializer

    queryset = GtmPool.objects.select_related("device").all()
    serializer_class = GtmPoolSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


# ---------------------------------------------------------------------------
# Firewall Policy ViewSets
# ---------------------------------------------------------------------------


class AddressBookViewSet(viewsets.ModelViewSet):
    from assets.models import AddressBook

    from .serializers import AddressBookSerializer

    queryset = AddressBook.objects.select_related("device", "parent").all()
    serializer_class = AddressBookSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class ServiceViewSet(viewsets.ModelViewSet):
    from assets.models import Service

    from .serializers import ServiceSerializer

    queryset = Service.objects.select_related("device").all()
    serializer_class = ServiceSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "protocol")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class PolicyViewSet(viewsets.ModelViewSet):
    from assets.models import Policy

    from .serializers import PolicySerializer

    queryset = Policy.objects.select_related("device").all()
    serializer_class = PolicySerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "policy_id", "device__hostname")
    ordering_fields = ("order", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class NatRuleViewSet(viewsets.ModelViewSet):
    from assets.models import NatRule

    from .serializers import NatRuleSerializer

    queryset = NatRule.objects.select_related("device").all()
    serializer_class = NatRuleSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "device__hostname")
    ordering_fields = ("order", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


# ---------------------------------------------------------------------------
# IPAM ViewSets
# ---------------------------------------------------------------------------


class TagViewSet(viewsets.ModelViewSet):
    from assets.models import Tag

    from .serializers import TagSerializer

    queryset = Tag.objects.all()
    serializer_class = TagSerializer
    permission_classes = (AllowAny,)
    search_fields = ("name",)
    ordering_fields = ("name", "created_at")


class SubnetViewSet(viewsets.ModelViewSet):
    from assets.models import Subnet

    from .serializers import SubnetSerializer

    queryset = Subnet.objects.prefetch_related("tags").all()
    serializer_class = SubnetSerializer
    permission_classes = (AllowAny,)
    search_fields = ("network", "description")
    ordering_fields = ("network", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        datacenter = self.request.query_params.get("datacenter")
        if datacenter:
            qs = qs.filter(datacenter_id=datacenter)
        tag = self.request.query_params.get("tag")
        if tag:
            qs = qs.filter(tags__id=tag)
        return qs


class IPAddressViewSet(viewsets.ModelViewSet):
    from assets.models import IPAddress

    from .serializers import IPAddressSerializer

    queryset = IPAddress.objects.select_related("subnet", "device", "security_zone").all()
    serializer_class = IPAddressSerializer
    permission_classes = (AllowAny,)
    search_fields = ("ip_address", "description", "device__hostname")
    ordering_fields = ("ip_address", "created_at")

    def get_queryset(self):
        qs = super().get_queryset()
        subnet = self.request.query_params.get("subnet")
        if subnet:
            qs = qs.filter(subnet_id=subnet)
        status = self.request.query_params.get("status")
        if status:
            qs = qs.filter(status=status)
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        return qs


class RouteViewSet(viewsets.ModelViewSet):
    from assets.models import Route

    from .serializers import RouteSerializer

    queryset = Route.objects.select_related("vrf", "vrf__device").all()
    serializer_class = RouteSerializer
    permission_classes = (AllowAny,)
    search_fields = ("destination", "nexthop", "vrf__name", "vrf__device__hostname")
    ordering_fields = ("destination", "protocol", "metric")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(vrf__device_id=device_id)
        protocol = self.request.query_params.get("protocol")
        if protocol:
            qs = qs.filter(protocol=protocol)
        return qs




class TopologyViewSet(viewsets.ModelViewSet):
    from assets.models import Topology

    from .serializers import TopologySerializer

    queryset = Topology.objects.all()
    serializer_class = TopologySerializer
    permission_classes = (AllowAny,)
    search_fields = ("name", "description")
    ordering_fields = ("name", "updated_at")


class ArpMacViewSet(viewsets.ModelViewSet):
    from assets.models import ArpMac

    from .serializers import ArpMacSerializer

    queryset = ArpMac.objects.select_related("device").all()
    serializer_class = ArpMacSerializer
    permission_classes = (AllowAny,)
    search_fields = ("ip_address", "mac_address", "device__hostname", "interface", "vendor")
    ordering_fields = ("ip_address", "mac_address", "vlan", "learned_at", "updated_at")

    def get_queryset(self):
        qs = super().get_queryset()
        device_id = self.request.query_params.get("device")
        if device_id:
            qs = qs.filter(device_id=device_id)
        vlan = self.request.query_params.get("vlan")
        if vlan:
            qs = qs.filter(vlan=vlan)
        arp_type = self.request.query_params.get("arp_type")
        if arp_type:
            qs = qs.filter(arp_type=arp_type)
        return qs


class SubnetUsageLogViewSet(viewsets.ModelViewSet):
    from assets.models import SubnetUsageLog

    from .serializers import SubnetUsageLogSerializer

    queryset = SubnetUsageLog.objects.select_related("subnet").all()
    serializer_class = SubnetUsageLogSerializer
    permission_classes = (AllowAny,)
    ordering_fields = ("recorded_at", "utilization")

    def get_queryset(self):
        qs = super().get_queryset()
        subnet_id = self.request.query_params.get("subnet")
        if subnet_id:
            qs = qs.filter(subnet_id=subnet_id)
        return qs
